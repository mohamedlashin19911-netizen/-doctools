"""توليد ملفات PowerPoint (PPTX) من مقاطع نصية وصور."""
import io

from pptx import Presentation
from pptx.util import Inches, Pt
from PIL import Image


def _normalize_png(raw: bytes) -> bytes:
    try:
        img = Image.open(io.BytesIO(raw))
        out = io.BytesIO()
        img.convert("RGB").save(out, "PNG")
        return out.getvalue()
    except Exception:
        return raw


def build_pptx(sections: list) -> bytes:
    """sections: [{'title': str, 'bullets': [str], 'images': [bytes]}]"""
    if not sections:
        sections = [{"title": "", "bullets": [], "images": []}]
    prs = Presentation()
    W, H = prs.slide_width, prs.slide_height
    for sec in sections:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        title = (sec.get("title") or "").strip()
        bullets = [b for b in sec.get("bullets", []) if str(b).strip()]
        images = sec.get("images", [])
        text_area_h = H - Inches(1.2)
        if images:
            img = Image.open(io.BytesIO(_normalize_png(images[0])))
            ratio = img.width / img.height
            w = min(Inches(7.5), Inches(9.5))
            h = int(w / ratio)
            max_h = int(H * 0.55)
            if h > max_h:
                h = max_h
                w = int(h * ratio)
            slide.shapes.add_picture(io.BytesIO(_normalize_png(images[0])),
                                     int((W - w) / 2), Inches(1.1), width=w, height=h)
            if title or bullets:
                tb = slide.shapes.add_textbox(Inches(0.4), Inches(0.15), W - Inches(0.8), Inches(0.9))
                p = tb.text_frame.paragraphs[0]
                p.text = title or bullets[0]
                p.runs[0].font.size = Pt(24)
                p.runs[0].font.bold = True
                bullets = []
        else:
            tb = slide.shapes.add_textbox(Inches(0.4), Inches(0.15), W - Inches(0.8), Inches(1.0))
            p = tb.text_frame.paragraphs[0]
            p.text = title
            p.runs[0].font.size = Pt(34)
            p.runs[0].font.bold = True
        if bullets:
            top = Inches(1.3) if not images else Inches(5.4)
            tb2 = slide.shapes.add_textbox(Inches(0.6), top, W - Inches(1.2),
                                           H - top - Inches(0.4))
            tf = tb2.text_frame
            tf.word_wrap = True
            for i, b in enumerate(bullets[:20]):
                par = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                par.text = str(b)
                if par.runs:
                    par.runs[0].font.size = Pt(17)
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def sections_from_text(text: str) -> list:
    """نص/Markdown → شرائح: عناوين # تقسم الشرائح، والباقي نقاط."""
    sections = []
    cur = None
    for line in text.splitlines():
        l = line.strip()
        if not l:
            continue
        m = __import__("re").match(r"^#{1,2}\s+(.*)", l)
        if m:
            if cur:
                sections.append(cur)
            cur = {"title": m.group(1).strip(), "bullets": []}
        elif cur is None:
            cur = {"title": l, "bullets": []}
        else:
            cur["bullets"].append(__import__("re").sub(r"^[-*•]\s*", "", l))
    if cur:
        sections.append(cur)
    if not sections:
        sections = [{"title": "مستند", "bullets": []}]
    return sections

"""الخادم المحلي — يستمع على 127.0.0.1 فقط، لا يمكن الوصول إليه من الشبكة."""
import io
import os
import re
import base64
from html import escape
from urllib.parse import quote

import pymupdf
from fastapi import FastAPI, UploadFile, File, Form, Request
from fastapi.responses import Response, JSONResponse, FileResponse
from fastapi.staticfiles import StaticFiles

from . import convert as C
from . import archive as A
from . import docgen
from . import ocr as O
from . import folderanalyze as F

app = FastAPI(title="أدوات المستندات")

MAX_SIZE = 100 * 1024 * 1024


def _read(file: UploadFile) -> tuple:
    data = file.file.read()
    if len(data) > MAX_SIZE:
        raise ValueError("الملف أكبر من 100MB")
    fname = file.filename or "file"
    stem = fname
    for suf in (".tar.gz", ".tar.bz2", ".tar.xz", ".tgz"):
        if fname.lower().endswith(suf):
            stem = fname[:-len(suf)]
            break
    else:
        stem = re.sub(r"\.[^.]+$", "", fname)
    stem = re.sub(r'[\\/:*?"<>|]', "_", stem)
    return stem, C.sniff(data, C.detect_ext(fname)), data


def _file_response(data: bytes, ext: str, name: str) -> Response:
    safe = quote(f"{name}.{ext}")
    return Response(
        content=data,
        media_type=C.MIME.get(ext, "application/octet-stream"),
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{safe}"},
    )


FORMAT_ORDER = [
    "pdf", "docx", "pptx", "xlsx", "xls", "csv", "tsv", "epub", "fb2", "cbz", "xps",
    "txt", "md", "html", "rtf", "odt",
    "png", "jpg", "jpeg", "webp", "gif", "bmp", "tiff", "ico", "avif", "heic", "heif",
    "jp2", "tga", "pcx", "ppm", "dds", "sgi", "im", "eps", "xbm", "svg",
    "cur", "qoi", "psd", "wmf", "icns", "mpo", "xpm",
    "json", "xml", "yaml", "toml", "ini", "cfg", "conf", "tex", "rst", "srt", "vtt", "log",
    "extract", "zip", "7z", "tar", "tar.gz",
]


@app.post("/api/formats")
async def formats(request: Request):
    body = await request.json()
    srcs = body.get("srcs", [])
    target_sets = [set(C.TARGETS.get(C.TEXT_ALIASES.get(s, s), [])) for s in srcs if s]
    if not target_sets:
        return {"targets": []}
    common = target_sets[0]
    for t in target_sets[1:]:
        common &= t
    return {"targets": [t for t in FORMAT_ORDER if t in common]}


@app.post("/api/inspect")
async def inspect_api(files: list[UploadFile] = File(...)):
    """يكشف النوع الحقيقي لكل ملف من محتواه ويعيد صيغ التحويل المتاحة."""
    out = []
    for f in files:
        try:
            stem, src, data = _read(f)
            out.append({"name": f.filename, "src": src or None,
                        "targets": [t for t in FORMAT_ORDER if t in C.TARGETS.get(src, [])]})
        except Exception:
            out.append({"name": f.filename, "src": None, "targets": []})
    return {"files": out}


TARGET_LABELS = {"extract": "فك الضغط", "tar.gz": "TAR.GZ", "heic": "HEIC", "heif": "HEIF"}


@app.post("/api/convert")
async def convert_api(files: list[UploadFile] = File(...), target: str = Form(...)):
    try:
        parsed = [_read(f) for f in files]
        if target in C.ARCHIVE_EXTS and not all(s in C.ARCHIVE_EXTS for _, s, _ in parsed):
            packed = []
            for stem, src, data in parsed:
                if src in C.ARCHIVE_EXTS:
                    packed.extend(A.read_archive(data, src))
                else:
                    packed.append((f"{stem}.{src or 'bin'}", data))
            return _file_response(A.write_archive(packed, target), target, "أرشيف")
        if all(s in C.ARCHIVE_EXTS for _, s, _ in parsed):
            outs, errs = [], []
            for stem, src, data in parsed:
                members = A.read_archive(data, src)
                if not members:
                    raise ValueError("الأرشيف فاضي")
                if target == "extract" or target == src:
                    outs.extend(members)
                elif target in C.ARCHIVE_EXTS:
                    outs.append((f"{stem}.{target}", A.write_archive(members, target)))
                else:
                    for name, mdata in members:
                        msrc = C.detect_ext(name)
                        mstem = re.sub(r"\.[^.]+$", "", name)
                        if msrc in C.ARCHIVE_EXTS:
                            errs.append(f"{name}: أرشيف داخلي")
                            continue
                        try:
                            r = C.convert_one(mdata, msrc, target, mstem)
                            outs.extend([(f"{mstem}.{r[2]}", r[1])] if r[0] == "file" else r[1])
                        except Exception as e:
                            errs.append(f"{name}: {e}")
            if not outs:
                raise ValueError("فشلت كل التحويلات — " + "؛ ".join(errs[:3]))
            if len(outs) == 1:
                name, data = outs[0]
                ext = C.detect_ext(name) or target
                return _file_response(data, ext, re.sub(r"\.[^.]+$", "", name))
            return _file_response(A.write_archive(outs, "zip"), "zip", "ملفات-محولة")
        result = C.convert_many(parsed, target)
        if result[0] == "zip":
            _, zdata = result[1][0]
            return _file_response(zdata, "zip", "ملفات-محولة")
        _, data, ext = result
        return _file_response(data, ext, parsed[0][0] if len(parsed) == 1 else "محول")
    except Exception as e:
        return JSONResponse({"error": str(e)}, status_code=400)


def _pdf_blocks(data: bytes):
    doc = C._open_pdf(data)
    blocks = []
    for page in doc:
        for p in page.get_text("blocks"):
            text = p[4].strip()
            if text:
                blocks.append({"type": "text", "html": f"<p>{text}</p>"})
    doc.close()
    return blocks


def _docx_blocks(data: bytes):
    import docx
    d = docx.Document(io.BytesIO(data))
    blocks = []
    for p in d.paragraphs:
        if not p.text.strip():
            continue
        style = (p.style.name or "").lower()
        m = re.search(r"(\d)", style) if style.startswith(("heading", "عنوان")) else None
        tag = f"h{m.group(1)}" if m else "p"
        blocks.append({"type": "text", "html": f"<{tag}>{escape(p.text)}</{tag}>"})
    for rid, rel in d.part.rels.items():
        if "image" in rel.reltype:
            raw = rel.target_part.blob
            img = Image_open(raw)
            blocks.append({"type": "image", "data": f"data:image/png;base64,{base64.b64encode(img).decode()}"})
    return blocks


def Image_open(raw: bytes) -> bytes:
    from PIL import Image
    try:
        img = Image.open(io.BytesIO(raw))
        out = io.BytesIO()
        img.convert("RGB").save(out, "PNG")
        return out.getvalue()
    except Exception:
        return raw


def _pptx_blocks(data: bytes):
    from pptx import Presentation
    prs = Presentation(io.BytesIO(data))
    blocks = []
    for slide in prs.slides:
        for sh in slide.shapes:
            if sh.shape_type == 13:
                try:
                    png = Image_open(sh.image.blob)
                    b64 = base64.b64encode(png).decode()
                    blocks.append({"type": "image", "data": f"data:image/png;base64,{b64}"})
                except Exception:
                    pass
            elif sh.has_text_frame:
                for para in sh.text_frame.paragraphs:
                    t = "".join(r.text for r in para.runs).strip()
                    if t:
                        tag = "h2" if (para.font.size and para.font.size.pt >= 24 or para.font.bold) else "p"
                        blocks.append({"type": "text", "html": f"<{tag}>{escape(t)}</{tag}>"})
    return blocks


@app.post("/api/extract")
async def extract(file: UploadFile = File(...)):
    """تحويل ملف مرفوع إلى blocks للمحرر (صور/نص/PDF/Word)."""
    try:
        stem, src, data = _read(file)
        if src in C.IMAGE_EXTS:
            norm = C._norm_image(data, "png") if src not in ("png", "jpg", "jpeg") else data
            b64 = base64.b64encode(norm).decode()
            return {"blocks": [{"type": "image", "data": f"data:image/png;base64,{b64}"}]}
        if src == "pdf":
            return {"blocks": _pdf_blocks(data)}
        if src == "docx":
            return {"blocks": _docx_blocks(data)}
        if src == "pptx":
            return {"blocks": _pptx_blocks(data)}
        if src == "svg":
            _, png, _ = C.svg_to(data, "png", "s")
            b64 = base64.b64encode(png).decode()
            return {"blocks": [{"type": "image", "data": f"data:image/png;base64,{b64}"}]}
        if src in ("xlsx", "csv", "xls"):
            rows = C._read_xlsx(data) if src == "xlsx" else (
                C._read_xls(data) if src == "xls" else C._read_csv(data))
            html = C.rows_to_html(rows)
            return {"blocks": [{"type": "text", "html": html}]}
        if src in ("txt", "md", "html", "rtf", "odt", "json", "xml", "log"):
            text = C._source_to_text(data, src)
            paras = [p for p in re.split(r"\n\s*\n", text) if p.strip()] or [text]
            blocks = []
            for p in paras:
                m = re.match(r"^(#{1,3})\s+(.*)", p, re.DOTALL)
                if m:
                    tag = f"h{len(m.group(1))}"
                    blocks.append({"type": "text", "html": f"<{tag}>{escape(m.group(2))}</{tag}>"})
                else:
                    blocks.append({"type": "text", "html": f"<p>{escape(p)}</p>"})
            return {"blocks": blocks}
        return JSONResponse({"error": f"لا يمكن إدراج صيغة {src} في المحرر"}, status_code=400)
    except Exception as e:
        return JSONResponse({"error": str(e)}, status_code=400)


CREATE_FORMATS = {"pdf", "docx", "png", "jpg", "html", "txt", "md", "pptx"}


@app.post("/api/create")
async def create(request: Request):
    try:
        body = await request.json()
        fmt = body.get("format", "pdf")
        if fmt not in CREATE_FORMATS:
            raise ValueError(f"صيغة غير مدعومة: {fmt}")
        name = re.sub(r'[\\/:*?"<>|]', "_", body.get("name", "مستند").strip()) or "مستند"
        blocks = body.get("blocks", [])
        if not blocks:
            raise ValueError("المحرر فاضي — أضف نص أو صورة الأول")
        data, ext = docgen.build(fmt, blocks, name)
        return _file_response(data, ext, name)
    except Exception as e:
        return JSONResponse({"error": str(e)}, status_code=400)


@app.get("/api/health")
async def health():
    return {"ok": True, "local_only": True}


@app.post("/api/folder")
async def folder_api(request: Request):
    """تحليل مجلد محلي كامل — شجرة + إحصاءات + فحص عمق لكل الطبقات."""
    try:
        body = await request.json()
        path = body.get("path", "").strip().strip('"')
        deep = bool(body.get("deep", True))
        if not path:
            raise ValueError("اكتب مسار المجلد الأول")
        return F.analyze(path, deep)
    except ValueError as e:
        return JSONResponse({"error": str(e)}, status_code=400)
    except Exception as e:
        return JSONResponse({"error": f"فشل التحليل: {e}"}, status_code=400)


@app.post("/api/folder/scan-file")
async def scan_file_api(request: Request):
    """فحص ملف واحد بالاسم الكامل — هل امتداده بيقول الحقيقة؟"""
    try:
        body = await request.json()
        p = body.get("path", "").strip().strip('"')
        if not os.path.isfile(p):
            raise ValueError("الملف مش موجود — اتأكد من المسار")
        size = os.path.getsize(p)
        data = F._read_magic(p, size)
        return {"path": p, "real": F.real_type(data), "size_h": F._human(size),
                "msg": F.inspect_file(p, size) or "لا توجد مشاكل — الملف سليم"}
    except ValueError as e:
        return JSONResponse({"error": str(e)}, status_code=400)
    except Exception as e:
        return JSONResponse({"error": f"فشل الفحص: {e}"}, status_code=400)


@app.post("/api/install-desktop")
def install_desktop():
    """ينشئ اختصار DocTools في قائمة Start على الجهاز الحالي (ويندوز فقط)."""
    import subprocess
    import sys
    import tempfile
    if os.name != "nt":
        return JSONResponse({"error": "التثبيت كتطبيق متاح على ويندوز فقط"}, status_code=400)
    root = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    dsk = os.path.join(root, "desktop.pyw")
    if not os.path.exists(dsk):
        return JSONResponse({"error": "ملف desktop.pyw غير موجود — نزّل النسخة الكاملة"}, status_code=400)
    pyw = os.path.join(os.path.dirname(sys.executable), "pythonw.exe")
    if not os.path.exists(pyw):
        pyw = sys.executable
    ico = os.path.join(root, "app.ico")
    icon_line = f"$lnk.IconLocation = '{ico},0'" if os.path.exists(ico) else ""
    ps = f"""
$ws = New-Object -ComObject WScript.Shell
$path = Join-Path $env:APPDATA 'Microsoft\\Windows\\Start Menu\\Programs\\DocTools.lnk'
$lnk = $ws.CreateShortcut($path)
$lnk.TargetPath = '{pyw}'
$lnk.Arguments = '"{dsk}"'
$lnk.WorkingDirectory = '{root}'
{icon_line}
$lnk.Description = 'Local document tools - 100% offline'
$lnk.Save()
if (Test-Path $path) {{ Write-Output 'OK' }} else {{ Write-Output 'MISSING' }}
"""
    tmp = os.path.join(tempfile.gettempdir(), "doctools_install.ps1")
    with open(tmp, "w", encoding="utf-16") as f:
        f.write(ps)
    try:
        r = subprocess.run(["powershell", "-NoProfile", "-ExecutionPolicy", "Bypass", "-File", tmp],
                           capture_output=True, text=True, timeout=30, creationflags=0x08000000)
        if "OK" in (r.stdout or ""):
            return {"ok": True}
        return JSONResponse({"error": "تعذر إنشاء الاختصار في قائمة Start"}, status_code=500)
    except Exception as e:
        return JSONResponse({"error": f"فشل التثبيت: {e}"}, status_code=500)


@app.post("/api/scan")
async def scan_api(files: list[UploadFile] = File(...)):
    """فحص دفعات من الملفات المرفوعة (شرائح أول 64KB لكل ملف) — اسم + نوع حقيقي."""
    out = []
    for i, f in enumerate(files):
        try:
            data = await f.read()
            name = f.filename or f"file{i}"
            rt = F.real_type(data[:65536])
            out.append({"i": i, "name": name, "real": rt, "msg": F.inspect_bytes(name, data[:65536])})
        except Exception as e:
            out.append({"i": i, "name": f.filename or "", "real": "", "msg": f"تعذر الفحص: {e}"})
    return {"results": out}


@app.post("/api/scan-archive")
async def scan_archive_api(file: UploadFile = File(...)):
    """فحص جوه أرشيف مرفوع بالكامل."""
    try:
        data = await file.read()
        if len(data) > F.MAX_ARC_SIZE:
            raise ValueError("الأرشيف أكبر من 300MB — مش بيتفحص من المتصفح")
        return F.archive_scan(file.filename or "archive", data)
    except ValueError as e:
        return JSONResponse({"error": str(e)}, status_code=400)
    except Exception as e:
        return JSONResponse({"error": f"فشل فحص الأرشيف: {e}"}, status_code=400)


@app.post("/api/ocr")
async def ocr_api(file: UploadFile = File(...), lang: str = Form("auto"), fmt: str = Form("simple")):
    """تفريغ نص من صورة/PDF/Word باستخدام محرك OCR المدمج في ويندوز."""
    try:
        stem, src, data = _read(file)
        if src == "pdf":
            text, engine = O.pdf_to_text(data, lang)
        elif src == "docx":
            text, engine = O.docx_to_text(data)
        elif src in C.IMAGE_EXTS:
            text, engine = O.image_to_text(data, lang)
        elif src in C.TEXT_LIKE or src in C.TABLE_EXTS:
            raise ValueError("الملفات النصية والجداول مش محتاجة تفريغ — استخدم المحوّل أو المحرر")
        else:
            raise ValueError("الصيغة لا تصلح للتفريغ: استخدم صورة أو PDF أو Word")
        if not text:
            raise ValueError("لم يتم العثور على نص في الملف")
        text = O.format_text(text, fmt if fmt in O.FORMAT_MODES else "simple")
        return {"text": text, "engine": engine, "src": src}
    except ValueError as e:
        return JSONResponse({"error": str(e)}, status_code=400)
    except Exception as e:
        return JSONResponse({"error": f"فشل التفريغ: {e}"}, status_code=400)


app.mount("/", StaticFiles(directory="static", html=True), name="static")
